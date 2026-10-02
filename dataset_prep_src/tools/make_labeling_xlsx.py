"""
tools/make_labeling_xlsx.py
===========================
Tạo file xlsx để nhóm labeling điền trực tiếp (đăng lên Google Drive/Sheets).

Đọc ``manifest.csv`` -> tạo xlsx gồm:
    * sheet "labels"     : ``image`` (điền sẵn) + các cột cần điền (plate_text +
                           8 cột toạ độ 4 góc) + ``track_id``/``frame_index`` (đối chiếu).
    * sheet "Hướng dẫn"  : tóm tắt cách điền.

Chạy::

    python tools/make_labeling_xlsx.py --manifest labeling/manifest.csv --output labeling/labeling.xlsx
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


#: Các cột người dán nhãn cần ĐIỀN (theo đúng thứ tự hiển thị).
FILLABLE = ["plate_text", "tl_x", "tl_y", "tr_x", "tr_y", "br_x", "br_y", "bl_x", "bl_y"]

HEADER_FILL = PatternFill("solid", fgColor="FFC000")   # vàng - tiêu đề
FILL_FILL = PatternFill("solid", fgColor="E2EFDA")     # xanh nhạt - cần điền
READONLY_FILL = PatternFill("solid", fgColor="EDEDED")  # xám - chỉ đọc


def main() -> int:
    ap = argparse.ArgumentParser(description="Tạo xlsx gán nhãn biển số từ manifest.csv.")
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--output", default="labeling/labeling.xlsx")
    args = ap.parse_args()

    rows = []
    with open(args.manifest, newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            rows.append(row)
    if not rows:
        raise SystemExit("manifest.csv rỗng.")

    wb = Workbook()
    ws = wb.active
    ws.title = "labels"

    headers = ["image"] + FILLABLE + ["track_id", "frame_index"]
    ws.append(headers)

    # Style tiêu đề.
    for cell in ws[1]:
        cell.font = Font(bold=True, color="000000")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Điền dữ liệu: image + track_id + frame_index có sẵn; các cột điền để trống.
    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=r["image"])
        ws.cell(row=i, column=11, value=int(r.get("track_id", 0)))
        ws.cell(row=i, column=12, value=int(r.get("frame_index", 0)))

    # Tô màu: cột 2..10 (plate_text + 8 toạ độ) = xanh nhạt (cần điền).
    # Cột 1, 11, 12 = xám (chỉ đọc).
    for i in range(2, len(rows) + 2):
        for col in range(2, 11):
            ws.cell(row=i, column=col).fill = FILL_FILL
        ws.cell(row=i, column=1).fill = READONLY_FILL
        ws.cell(row=i, column=11).fill = READONLY_FILL
        ws.cell(row=i, column=12).fill = READONLY_FILL

    # Đóng băng dòng tiêu đề + cột image.
    ws.freeze_panes = "B2"

    # Độ rộng cột.
    ws.column_dimensions["A"].width = 12
    for col in range(2, 11):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.column_dimensions["K"].width = 10
    ws.column_dimensions["L"].width = 12

    # Sheet hướng dẫn.
    guide = wb.create_sheet("Hướng dẫn")
    guide_lines = [
        "HƯỚNG DẪN ĐIỀN NHÃN BIỂN SỐ",
        "",
        "1. Mỗi dòng = 1 ảnh ở cột 'image' (ảnh nằm trong thư mục images/ của zip).",
        "2. Cột plate_text: gõ biển số thật, bỏ dấu gạch/dấu chấm/khoảng trắng, CHỈ gồm 0-9 và A-Z.",
        "   Ví dụ: 51F-12345 -> 51F12345 ; 29-B1 12345 -> 29B112345.",
        "   Không đọc nổi text -> ghi 'unreadable'.",
        "3. 8 cột toạ độ là 4 GÓC của TẤM BIỂN, theo thứ tự chiều kim đồng hồ:",
        "   TL = góc trên-trái   -> (tl_x, tl_y)",
        "   TR = góc trên-phải   -> (tr_x, tr_y)",
        "   BR = góc dưới-phải   -> (br_x, br_y)",
        "   BL = góc dưới-trái   -> (bl_x, bl_y)",
        "   Đơn vị: pixel; gốc toạ độ (0,0) = góc TRÊN-TRÁI của ảnh.",
        "",
        "4. THỨ TỰ 4 GÓC PHẢI ĐÚNG TL -> TR -> BR -> BL. Đây là yêu cầu quan trọng nhất.",
        "5. Đặt điểm đúng MÉP GÓC của tấm biển (không phải khung ảnh, không lệch vào chữ).",
        "6. Cột track_id, frame_index là thông tin đối chiếu — KHÔNG cần sửa.",
    ]
    for line in guide_lines:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    guide["A1"].font = Font(bold=True, size=14)

    wb.save(args.output)
    print(f"Created: {args.output} | rows={len(rows)} | sheets={wb.sheetnames}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
