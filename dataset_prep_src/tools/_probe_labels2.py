"""Probe tạm 2: kích thước ảnh crop + kiểm tra giả định percent -> pixel."""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from PIL import Image

DEP = Path(r"S:\M_AGENT\CV\dataset_prep_src")
L = Path(r"S:\M_AGENT\CV\labeling")
RES = L / "result"


def dims(name: str):
    p = L / name
    return Image.open(p).size if p.exists() else None


print("kich thuoc vai anh crop (w,h):")
for n in ["000001.jpg", "000401.jpg", "000407.jpg", "000600.jpg", "000601.jpg",
          "000605.jpg", "000800.jpg", "000801.jpg", "001000.jpg", "001700.jpg"]:
    print(f"  {n}: {dims(n)}")

# --- 401-600: original_width co khop kich thuoc file? ------------------
rows = list(csv.DictReader((RES / "401-600.csv").open(encoding="utf-8")))
mismatch = 0
npoints = Counter()
samples = []
for r in rows:
    name = r["image"].split("%5C")[-1]
    obj = json.loads(r["label"]) if r["label"].strip() else []
    if not obj:
        continue
    item = obj[0]
    w, h = dims(name) or (None, None)
    npoints[len(item["points"])] += 1
    if w is not None and (item["original_width"] != w or item["original_height"] != h):
        mismatch += 1
        if len(samples) < 3:
            samples.append((name, item["original_width"], item["original_height"], w, h))
print(f"\n401-600: so diem trong polygon: {dict(npoints)}")
print(f"401-600: so dong original_w/h KHAC kich thuoc file: {mismatch}")
for s in samples:
    print("   ", s)

# --- text chuan hoa + co polygon hay khong ----------------------------
def norm(t: str) -> str:
    t = (t or "").strip()
    return "no_plate" if t.lower() in ("no plate", "no_plate") else t


for f, col in [("labels_anh_601_800.csv", "plate_text")]:
    rows2 = list(csv.DictReader((RES / f).open(encoding="utf-8")))
    c = Counter()
    for r in rows2:
        has = any(float(r[k]) != 0 for k in ("tl_x", "tl_y", "tr_x", "tr_y",
                                             "br_x", "br_y", "bl_x", "bl_y"))
        c[(norm(r[col]), has)] += 1
    print(f"\n{f}: (nhan, co_toa_do) -> {dict(c)}")

from openpyxl import load_workbook

wb = load_workbook(RES / "801-1701.xlsx", read_only=True, data_only=True)
ws = wb["Sheet1"]
it = ws.iter_rows(values_only=True)
next(it)
c = Counter()
for r in it:
    if not r or not r[0]:
        continue
    has = any(float(r[i] or 0) != 0 for i in range(2, 10))
    c[(norm(str(r[1])), has)] += 1
print(f"\n801-1701.xlsx: (nhan, co_toa_do) -> {dict(c)}")
